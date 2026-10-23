"""dev.60 文件轉 PDF：Impress 匯出選項、繁體中文字形、轉換後的頁數與原檔版面說明。

測試檔用 openpyxl／python-pptx 產生（沒有就略過）；實際轉換需要 LibreOffice（沒有就略過）。
執行：python3 -m unittest agent/tests/test_office.py -v
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)
from worker import office  # noqa: E402
import make_media  # noqa: E402

try:
    import openpyxl  # noqa: F401
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False
try:
    import pptx  # noqa: F401
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False
HAS_LO = bool(shutil.which("soffice") or shutil.which("libreoffice"))


class OfficeOptionsTest(unittest.TestCase):
    def test_impress_export_options_are_explicit(self):
        arg = office.convert_arg("PPTX")
        self.assertTrue(arg.startswith("pdf:impress_pdf_Export:"))
        opts = json.loads(arg.split(":", 2)[2])
        self.assertEqual({k: v["value"] for k, v in opts.items()},
                         {"ExportNotesPages": "false", "ExportOnlyNotesPages": "false", "ExportHiddenSlides": "false"})
        self.assertEqual(office.convert_arg("ppt")[:22], "pdf:impress_pdf_Export")
        for ext in ("docx", "xlsx", "doc", "xls", "txt"):
            self.assertEqual(office.convert_arg(ext), "pdf", "Word／Excel 照原檔的頁面設定，不加選項")

    def test_profile_sets_traditional_chinese_for_asian_text(self):
        d = tempfile.mkdtemp()
        try:
            office.seed_profile(d)
            with open(os.path.join(d, "user", "registrymodifications.xcu"), encoding="utf-8") as f:
                x = f.read()
            self.assertIn("DefaultLocale_CJK", x)
            self.assertIn("<value>zh-TW</value>", x)
        finally:
            shutil.rmtree(d, ignore_errors=True)

    def test_pdf_page_count(self):
        d = tempfile.mkdtemp()
        try:
            p = os.path.join(d, "a.pdf")
            make_media.pdf(p)
            self.assertEqual(office.pdf_pages(p), 1)
            self.assertIsNone(office.pdf_pages(os.path.join(d, "none.pdf")))
        finally:
            shutil.rmtree(d, ignore_errors=True)


@unittest.skipUnless(HAS_XLSX and HAS_PPTX, "需要 openpyxl 與 python-pptx 產生測試檔")
class LayoutNotesTest(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp(prefix="office-")
        make_media.office_demo(self.d)
        self.x = os.path.join(self.d, "OfficeDemo", "D02.xlsx")
        self.p = os.path.join(self.d, "OfficeDemo", "D03.pptx")

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def test_pptx_hidden_slides(self):
        self.assertEqual(office.pptx_slides(self.p), (4, 1))
        self.assertIsNone(office.pptx_slides(self.x), "不是簡報")

    def test_xlsx_points_out_original_print_settings(self):
        notes = office.xlsx_layout_notes(self.x)
        self.assertEqual(len(notes), 1, notes)
        self.assertIn("工作表「課程表」", notes[0])
        self.assertIn("A4 直式", notes[0])
        self.assertIn("調整成一頁寬", notes[0])
        self.assertFalse(any("「一頁寬」" in n for n in notes), "第 2 個工作表已設定調整成一頁寬：不提醒")

    def test_xlsx_print_area_and_hidden_sheet(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.x)
        wb["課程表"].print_area = "A1:F20"
        wb.create_sheet("備用").sheet_state = "hidden"
        wb.save(self.x)
        text = "\n".join(office.xlsx_layout_notes(self.x))
        self.assertIn("列印範圍 A1:F20", text)
        self.assertIn("工作表「備用」是隱藏的", text)
        self.assertNotIn("寬度約", text, "有列印範圍時以列印範圍為準")

    def test_broken_files_do_not_raise(self):
        bad = os.path.join(self.d, "bad.xlsx")
        with open(bad, "wb") as f:
            f.write(b"not a zip")
        self.assertEqual(office.xlsx_layout_notes(bad), [])
        self.assertIsNone(office.pptx_slides(bad))


@unittest.skipUnless(HAS_XLSX and HAS_PPTX and HAS_LO, "需要 LibreOffice")
class RealConversionTest(unittest.TestCase):
    """真的用 LibreOffice 轉換（與 Agent 映像同為 Ubuntu 24.04 的 LibreOffice 24.2）。"""

    @classmethod
    def setUpClass(cls):
        cls.d = tempfile.mkdtemp(prefix="office-real-")
        make_media.office_demo(cls.d)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.d, ignore_errors=True)

    def convert(self, name: str) -> tuple[str, dict]:
        src = os.path.join(self.d, "OfficeDemo", name)
        out = tempfile.mkdtemp(dir=self.d)
        prof = tempfile.mkdtemp(dir=self.d)
        pdf = office.convert_to_pdf(src, out, prof)
        return pdf, office.layout_check(src, os.path.splitext(name)[1][1:], pdf)

    def test_pptx_pages_equal_visible_slides(self):
        pdf, chk = self.convert("D03.pptx")
        self.assertEqual(chk["pages"], 3, "4 張投影片、1 張隱藏 → 3 頁；講者備註不會多出頁面")
        self.assertTrue(any("1 張隱藏" in n for n in chk["notes"]))
        self.assertFalse(any(n.startswith("⚠") for n in chk["notes"]), chk)

    def test_xlsx_uses_traditional_chinese_glyphs(self):
        pdf, chk = self.convert("D02.xlsx")
        with open(pdf, "rb") as f:
            data = f.read()
        self.assertIn(b"CJKtc", data, "中文用繁體中文字形（Noto CJK TC）")
        self.assertNotIn(b"CJKsc", data, "不用簡體字形")
        self.assertGreaterEqual(chk["pages"], 2)
        self.assertTrue(any("課程表" in n for n in chk["notes"]), chk)


if __name__ == "__main__":
    unittest.main(verbosity=2)
